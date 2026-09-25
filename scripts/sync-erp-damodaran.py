#!/usr/bin/env python3
"""ERP Indonesia - dataset negara Damodaran (NYU Stern). Sumber resmi akademik.

Berkas: ctrypremJuly<YY>.xlsx (edisi Juli) dan ctryprem.xlsx (edisi Januari).
Pilih edisi dengan 'Date of update' terbaru. Kolom diambil apa adanya; TIDAK
ada angka cadangan - kalau kolom/baris tidak ada, keluar GAGAL.
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import io
import json
import re
import sys
import urllib.request

BASE = "https://pages.stern.nyu.edu/~adamodar/pc/datasets/"
SHEET = "ERPs by country"
NEGARA = "Indonesia"
KOLOM_ERP = ("Total Equity Risk Premium", "Country Risk Premium", "Total Equity Risk Premium2")
BULAN = {"january": 1, "february": 2, "march": 3, "april": 4, "may": 5, "june": 6, "july": 7,
         "august": 8, "september": 9, "october": 10, "november": 11, "december": 12}


def gagal(reason: str, **extra) -> None:
    print(json.dumps({"status": "GAGAL", "reason": reason, **extra}, ensure_ascii=False))
    sys.exit(0)


def unduh(url: str) -> bytes | None:
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    try:
        with urllib.request.urlopen(req, timeout=60) as response:
            if response.status != 200:
                return None
            return response.read()
    except Exception:
        return None


def impor_openpyxl():
    try:
        import openpyxl
    except Exception:
        gagal("OPENPYXL_TIDAK_TERSEDIA")
    return openpyxl


def sel(baris: tuple) -> list:
    return [c.value for c in baris]


def teks(value) -> str:
    return "" if value is None else str(value).strip()


def normal(value) -> float | None:
    if isinstance(value, (int, float)):
        return float(value) * 100
    match = re.search(r"-?\d+(?:[.,]\d+)?", teks(value))
    return float(match.group(0).replace(",", ".")) if match else None


def tanggal_teks(value) -> str | None:
    if isinstance(value, (dt.datetime, dt.date)):
        return value.date().isoformat() if isinstance(value, dt.datetime) else value.isoformat()
    match = re.match(r"^([A-Za-z]+)\s+(\d{1,2}),\s*(\d{4})$", teks(value))
    if match and match.group(1).lower() in BULAN:
        return dt.date(int(match.group(3)), BULAN[match.group(1).lower()], int(match.group(2))).isoformat()
    return None


def baca_edisi(nama: str, konten: bytes) -> dict | None:
    openpyxl = impor_openpyxl()
    wb = openpyxl.load_workbook(io.BytesIO(konten), data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        return None
    ws = wb[SHEET]
    tanggal = None
    indeks = None
    mature = None
    baris_negara = None
    for baris in ws.iter_rows(min_row=1, max_row=ws.max_row):
        nilai = sel(baris)
        kolom_a = teks(nilai[0]).lower() if nilai else ""
        if kolom_a == "date of update:" and len(nilai) > 1 and tanggal is None:
            tanggal = tanggal_teks(nilai[1])
        if kolom_a.startswith("enter the current risk premium") and len(nilai) > 4 and mature is None:
            mature = normal(nilai[4])
        if indeks is None and any(teks(v).lower() == "total equity risk premium" for v in nilai):
            indeks = {teks(v).lower(): i for i, v in enumerate(nilai) if teks(v)}
            continue
        if kolom_a == NEGARA.lower() and indeks is not None:
            baris_negara = nilai
            break
    if baris_negara is None or not indeks:
        return None
    hasil = {"nama_edisi": nama, "tanggal_update": tanggal, "mature_erp_pct": mature}
    for label in KOLOM_ERP:
        posisi = indeks.get(label.lower())
        hasil[label] = normal(baris_negara[posisi]) if posisi is not None and posisi < len(baris_negara) else None
    hasil["nilai_pct"] = hasil.get(KOLOM_ERP[0])
    hasil["sha256"] = hashlib.sha256(konten).hexdigest()
    hasil["url"] = BASE + nama
    return hasil


def kandidat_edisi(tahun: int) -> list[str]:
    return [f"ctrypremJuly{str(tahun)[2:]}.xlsx", "ctryprem.xlsx"]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output")
    args = parser.parse_args()

    tahun = dt.date.today().year
    edisi = []
    for nama in kandidat_edisi(tahun):
        konten = unduh(BASE + nama)
        if konten is None:
            continue
        hasil = baca_edisi(nama, konten)
        if hasil and hasil.get("nilai_pct") is not None:
            edisi.append(hasil)
    if not edisi:
        gagal("EDISI_DAMODARAN_TIDAK_TERBACA")

    terpilih = max(edisi, key=lambda item: item.get("tanggal_update") or "")
    if not terpilih.get("tanggal_update"):
        gagal("TANGGAL_UPDATE_TIDAK_TERBACA", edisi=[e["nama_edisi"] for e in edisi])

    nilai = terpilih["nilai_pct"]
    if not isinstance(nilai, float) or not 1.0 <= nilai <= 25.0:
        gagal("NILAI_ERP_DI_LUAR_AKAL", nilai=nilai)

    as_of = dt.date.fromisoformat(terpilih["tanggal_update"])
    if as_of > dt.date.today():
        gagal("TANGGAL_UPDATE_MASA_DEPAN", tanggal_update=terpilih["tanggal_update"])
    umur_hari = (dt.date.today() - as_of).days
    if umur_hari > 400:
        gagal("EDISI_TERLALU_TUA", umur_hari=umur_hari)

    keluaran = {
        "status": "SUKSES",
        "input_key": "EQUITY_RISK_PREMIUM_PCT",
        "nilai_pct": round(nilai, 4),
        "kolom": KOLOM_ERP[0],
        "edisi": terpilih["nama_edisi"],
        "tanggal_update": terpilih["tanggal_update"],
        "umur_hari": umur_hari,
        "mature_erp_pct": terpilih.get("mature_erp_pct"),
        "country_risk_premium_pct": terpilih.get(KOLOM_ERP[1]),
        "erp_cds_pct": terpilih.get(KOLOM_ERP[2]),
        "sumber_nama": "Aswath Damodaran (NYU Stern) - Country Risk Premiums",
        "sumber_url": terpilih["url"],
        "berkas_sha256": terpilih["sha256"],
        "negara": NEGARA,
        "edisi_dipertimbangkan": [e["nama_edisi"] for e in edisi],
    }
    teks_keluar = json.dumps(keluaran, ensure_ascii=False)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as handle:
            handle.write(teks_keluar)
    print(teks_keluar)


if __name__ == "__main__":
    main()